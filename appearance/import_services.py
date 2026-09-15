import csv
import hashlib
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .models import (
    AppearanceApprovalRecord,
    DataUpload,
    HiBobEmployee,
    ImportBatch,
    ImportIssue,
    SecurityInfoRecord,
    TattooRecord,
)
from .services import normalize_employee_id


TATTOO_HEADER_ALIASES = {
    "id": "employee_id",
    "name": "source_name",
    "tattoos": "tattoo_details",
    "date": "source_date",
    "shouldbecovered": "should_be_covered",
    "connotation": "connotation",
    "where": "location",
}

APPROVAL_HEADER_ALIASES = {
    "id": "employee_id",
    "name": "source_name",
    "responsable": "responsible",
    "responsible": "responsible",
    "situation": "situation",
    "testtimeframeneeded": "test_time_frame_needed",
    "medicalcondition": "medical_condition",
    "paperworksubmmited": "paperwork_submitted",
    "paperworksubmitted": "paperwork_submitted",
    "testinitialdate": "test_initial_date",
    "testfinaldate": "test_final_date",
    "status": "status",
    "coments": "comments",
    "comments": "comments",
}

SECURITY_HEADER_ALIASES = {
    "employeeid": "employee_id",
    "idempleado": "employee_id",
    "idcolaborador": "employee_id",
    "numeroempleado": "employee_id",
    "nroempleado": "employee_id",
    "documento": "document",
    "nombre": "first_name",
    "apellidos": "last_name",
    "numerocontacto": "contact_number",
    "rh": "blood_type",
    "tipo": "vehicle_type",
    "marca": "vehicle_brand",
    "modelo": "vehicle_model",
    "color": "vehicle_color",
    "placasvehiculo": "vehicle_plate",
    "departamento": "department",
    "cargo": "role",
    "notarjeta": "card_number",
    "notarjetasec": "secondary_card_number",
    "facilitycodewfm": "facility_code_wfm",
    "facilitycodese": "facility_code_se",
    "fotoydatos": "photo_and_data",
    "eps": "eps",
    "notarjeta6digitos": "six_digit_card_number",
    "reposicion": "first_replacement",
    "fecha1erareposicion": "first_replacement_date",
    "segundareposicion": "second_replacement",
    "fecha2dareposicion": "second_replacement_date",
}

SECURITY_REQUIRED_FIELDS = {
    "document",
    "first_name",
    "last_name",
    "department",
    "role",
    "card_number",
}

SUPPORTED_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


class ImportServiceError(Exception):
    pass


def normalize_header(value):
    value = str(value or "").strip().lower()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(character for character in value if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]", "", value)


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def file_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_yes_no(value):
    normalized = clean_text(value).lower()
    if not normalized:
        return None
    if normalized in {"yes", "y", "si", "sí", "true", "1"}:
        return True
    if normalized in {"no", "n", "false", "0"}:
        return False
    return None


def normalize_answer(value):
    normalized = clean_text(value).lower().replace(".", "")
    if normalized in {"yes", "y", "si", "sí", "true", "1"}:
        return AppearanceApprovalRecord.Answer.YES
    if normalized in {"no", "n", "false", "0"}:
        return AppearanceApprovalRecord.Answer.NO
    if normalized in {"na", "n/a", "not applicable"}:
        return AppearanceApprovalRecord.Answer.NA
    return AppearanceApprovalRecord.Answer.UNKNOWN


def normalize_approval_status(value):
    normalized = re.sub(r"\s+", " ", clean_text(value).lower())
    if normalized == "approved":
        return AppearanceApprovalRecord.ApprovalStatus.APPROVED
    if normalized == "rejected":
        return AppearanceApprovalRecord.ApprovalStatus.REJECTED
    if normalized in {"in progress", "inprogress"}:
        return AppearanceApprovalRecord.ApprovalStatus.IN_PROGRESS
    return AppearanceApprovalRecord.ApprovalStatus.UNKNOWN


def parse_optional_date(value, workbook_epoch):
    if value in {None, ""}:
        return None, ""

    if isinstance(value, datetime):
        parsed = value.date()
        return parsed, parsed.isoformat()
    if isinstance(value, date):
        return value, value.isoformat()
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, workbook_epoch)
            parsed = converted.date() if isinstance(converted, datetime) else converted
            return parsed, parsed.isoformat()
        except (ValueError, TypeError, OverflowError):
            return None, clean_text(value)

    raw = clean_text(value)
    if raw.lower().replace(".", "") in {"na", "n/a", "not applicable"}:
        return None, raw

    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).date(), raw
        except ValueError:
            continue
    return None, raw


def _build_header_map(row, aliases):
    header_map = {}
    for index, value in enumerate(row):
        mapped = aliases.get(normalize_header(value))
        if mapped:
            header_map[mapped] = index
    return header_map


def _build_security_header_map(row):
    header_map = _build_header_map(row, SECURITY_HEADER_ALIASES)

    if "employee_id" not in header_map and "document" in header_map:
        document_index = header_map["document"]
        previous_index = document_index - 1
        if previous_index >= 0 and not normalize_header(row[previous_index]):
            header_map["employee_id"] = previous_index

    return header_map


def _find_header_sheet(workbook, aliases, required_fields):
    candidates = []
    for sheet in workbook.worksheets:
        max_search_row = min(sheet.max_row or 1, 15)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=max_search_row, values_only=True),
            start=1,
        ):
            header_map = _build_header_map(row, aliases)
            if required_fields.issubset(header_map):
                candidates.append((sheet.max_row or 0, sheet, header_map, row_number))
                break

    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])


def find_tattoo_sheet(workbook):
    return _find_header_sheet(
        workbook,
        TATTOO_HEADER_ALIASES,
        {"employee_id", "tattoo_details", "should_be_covered"},
    )


def find_approval_sheet(workbook):
    return _find_header_sheet(
        workbook,
        APPROVAL_HEADER_ALIASES,
        {"employee_id", "situation", "status", "comments"},
    )


def find_security_sheet(workbook):
    candidates = []
    for sheet in workbook.worksheets:
        max_search_row = min(sheet.max_row or 1, 15)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=max_search_row, values_only=True),
            start=1,
        ):
            header_map = _build_security_header_map(row)
            if SECURITY_REQUIRED_FIELDS.issubset(header_map):
                candidates.append((sheet.max_row or 0, sheet, header_map, row_number))
                break

    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])


def find_security_csv_header(rows):
    for index, row in enumerate(rows[:15]):
        header_map = _build_security_header_map(row)
        if SECURITY_REQUIRED_FIELDS.issubset(header_map):
            return header_map, index
    return None


def detect_source_type(workbook):
    matches = []
    if find_tattoo_sheet(workbook):
        matches.append(ImportBatch.SourceType.TATTOOS)
    if find_approval_sheet(workbook):
        matches.append(ImportBatch.SourceType.APPEARANCE_APPROVALS)
    if find_security_sheet(workbook):
        matches.append(ImportBatch.SourceType.SECURITY_GENERAL)

    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ImportServiceError("Workbook matches more than one supported source format.")
    raise ImportServiceError(
        "Workbook format was not recognized. Expected the GP tattoo tracker, "
        "Appearance approvals tracker, or Security general base."
    )


def _known_hibob_ids(employee_ids):
    if not employee_ids:
        return set()
    return set(
        HiBobEmployee.objects.using("hibob")
        .filter(employee_id__in=employee_ids)
        .values_list("employee_id", flat=True)
        .distinct()
    )


def _issue(batch, row_number, employee_id, message):
    ImportIssue.objects.create(
        batch=batch,
        row_number=row_number,
        employee_id=employee_id,
        message=message,
    )


def _finish_batch(batch, imported, rejected):
    batch.rows_imported = imported
    batch.rows_rejected = rejected
    batch.status = ImportBatch.Status.PARTIAL if rejected or batch.issues.exists() else ImportBatch.Status.SUCCESS
    batch.finished_at = timezone.now()
    batch.save(
        update_fields=[
            "rows_imported",
            "rows_rejected",
            "status",
            "finished_at",
        ]
    )
    return batch


def _read_tattoo_rows(sheet, header_map, header_row, workbook_epoch):
    records = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        def get(field):
            index = header_map.get(field)
            return values[index] if index is not None and index < len(values) else None

        employee_id = normalize_employee_id(get("employee_id"))
        relevant_values = [get(field) for field in (
            "employee_id",
            "source_name",
            "tattoo_details",
            "source_date",
            "should_be_covered",
            "connotation",
            "location",
        )]
        if not employee_id and all(value in {None, ""} for value in relevant_values):
            continue

        source_date, _ = parse_optional_date(get("source_date"), workbook_epoch)
        records.append({
            "row_number": row_number,
            "employee_id": employee_id,
            "source_name": clean_text(get("source_name")),
            "tattoo_details": clean_text(get("tattoo_details")),
            "should_be_covered": normalize_yes_no(get("should_be_covered")),
            "connotation": clean_text(get("connotation")),
            "location": clean_text(get("location")),
            "source_date": source_date,
            "raw_cover_value": clean_text(get("should_be_covered")),
        })
    return records


def _import_tattoos(workbook, batch, source_file):
    candidate = find_tattoo_sheet(workbook)
    if candidate is None:
        raise ImportServiceError("Tattoo tracker sheet was not found.")

    _, sheet, header_map, header_row = candidate
    rows = _read_tattoo_rows(sheet, header_map, header_row, workbook.epoch)
    batch.rows_received = len(rows)
    batch.save(update_fields=["rows_received"])

    known_ids = _known_hibob_ids([row["employee_id"] for row in rows if row["employee_id"]])
    valid_rows = []
    rejected_rows = 0

    for row in rows:
        if not row["employee_id"]:
            rejected_rows += 1
            _issue(batch, row["row_number"], "", "Missing Employee ID.")
            continue
        if row["employee_id"] not in known_ids:
            rejected_rows += 1
            _issue(batch, row["row_number"], row["employee_id"], "Employee ID was not found in HiBob.")
            continue
        if row["raw_cover_value"] and row["should_be_covered"] is None:
            _issue(
                batch,
                row["row_number"],
                row["employee_id"],
                f"Unrecognized SHOULD BE COVERED? value: {row['raw_cover_value']!r}. Imported as not specified.",
            )
        valid_rows.append(row)

    if rows and not valid_rows:
        raise ImportServiceError("No tattoo rows matched HiBob. Import aborted; the previous active snapshot was kept.")

    with transaction.atomic(using="default"):
        TattooRecord.objects.filter(is_active=True).update(is_active=False)
        TattooRecord.objects.bulk_create([
            TattooRecord(
                employee_id=row["employee_id"],
                source_name=row["source_name"],
                tattoo_details=row["tattoo_details"],
                should_be_covered=row["should_be_covered"],
                connotation=row["connotation"],
                location=row["location"],
                source_date=row["source_date"],
                source_sheet=sheet.title,
                source_file=source_file,
                import_batch=batch,
            )
            for row in valid_rows
        ])

    return _finish_batch(batch, len(valid_rows), rejected_rows)


def _read_approval_rows(sheet, header_map, header_row, workbook_epoch):
    records = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        def get(field):
            index = header_map.get(field)
            return values[index] if index is not None and index < len(values) else None

        employee_id = normalize_employee_id(get("employee_id"))
        relevant_values = [get(field) for field in set(APPROVAL_HEADER_ALIASES.values())]
        if not employee_id and all(value in {None, ""} for value in relevant_values):
            continue

        initial_date, initial_raw = parse_optional_date(get("test_initial_date"), workbook_epoch)
        final_date, final_raw = parse_optional_date(get("test_final_date"), workbook_epoch)

        records.append({
            "row_number": row_number,
            "employee_id": employee_id,
            "source_name": clean_text(get("source_name")),
            "responsible": clean_text(get("responsible")),
            "situation": clean_text(get("situation")),
            "test_time_frame_needed": normalize_answer(get("test_time_frame_needed")),
            "medical_condition": normalize_answer(get("medical_condition")),
            "paperwork_submitted": normalize_answer(get("paperwork_submitted")),
            "test_initial_date": initial_date,
            "test_initial_date_raw": initial_raw,
            "test_final_date": final_date,
            "test_final_date_raw": final_raw,
            "status": normalize_approval_status(get("status")),
            "status_raw": clean_text(get("status")),
            "comments": clean_text(get("comments")),
        })
    return records


def _import_approvals(workbook, batch, source_file):
    candidate = find_approval_sheet(workbook)
    if candidate is None:
        raise ImportServiceError("Appearance approvals tracker sheet was not found.")

    _, sheet, header_map, header_row = candidate
    rows = _read_approval_rows(sheet, header_map, header_row, workbook.epoch)
    batch.rows_received = len(rows)
    batch.save(update_fields=["rows_received"])

    known_ids = _known_hibob_ids([row["employee_id"] for row in rows if row["employee_id"]])
    valid_rows = []
    rejected_rows = 0

    for row in rows:
        if not row["employee_id"]:
            rejected_rows += 1
            _issue(batch, row["row_number"], "", "Missing Employee ID.")
            continue
        if row["employee_id"] not in known_ids:
            rejected_rows += 1
            _issue(batch, row["row_number"], row["employee_id"], "Employee ID was not found in HiBob.")
            continue
        if row["status"] == AppearanceApprovalRecord.ApprovalStatus.UNKNOWN and row["status_raw"]:
            _issue(
                batch,
                row["row_number"],
                row["employee_id"],
                f"Unrecognized Status value: {row['status_raw']!r}. Imported as not specified.",
            )
        if row["test_initial_date_raw"] and row["test_initial_date"] is None and row["test_initial_date_raw"].lower().replace(".", "") not in {"na", "n/a", "not applicable"}:
            _issue(
                batch,
                row["row_number"],
                row["employee_id"],
                f"Test initial date could not be parsed: {row['test_initial_date_raw']!r}. Raw value was preserved.",
            )
        if row["test_final_date_raw"] and row["test_final_date"] is None and row["test_final_date_raw"].lower().replace(".", "") not in {"na", "n/a", "not applicable"}:
            _issue(
                batch,
                row["row_number"],
                row["employee_id"],
                f"Test final date could not be parsed: {row['test_final_date_raw']!r}. Raw value was preserved.",
            )
        valid_rows.append(row)

    if rows and not valid_rows:
        raise ImportServiceError("No appearance approval rows matched HiBob. Import aborted; the previous active snapshot was kept.")

    with transaction.atomic(using="default"):
        AppearanceApprovalRecord.objects.filter(is_active=True).update(is_active=False)
        AppearanceApprovalRecord.objects.bulk_create([
            AppearanceApprovalRecord(
                employee_id=row["employee_id"],
                source_name=row["source_name"],
                responsible=row["responsible"],
                situation=row["situation"],
                test_time_frame_needed=row["test_time_frame_needed"],
                medical_condition=row["medical_condition"],
                paperwork_submitted=row["paperwork_submitted"],
                test_initial_date=row["test_initial_date"],
                test_initial_date_raw=row["test_initial_date_raw"],
                test_final_date=row["test_final_date"],
                test_final_date_raw=row["test_final_date_raw"],
                status=row["status"],
                comments=row["comments"],
                source_sheet=sheet.title,
                source_file=source_file,
                import_batch=batch,
            )
            for row in valid_rows
        ])

    return _finish_batch(batch, len(valid_rows), rejected_rows)


def _read_security_rows(values_iter, header_map, first_row_number):
    records = []
    security_fields = set(SECURITY_HEADER_ALIASES.values()) | {"employee_id"}

    for row_number, values in enumerate(values_iter, start=first_row_number):
        def get(field):
            index = header_map.get(field)
            return values[index] if index is not None and index < len(values) else None

        relevant_values = [get(field) for field in security_fields]
        if all(clean_text(value) == "" for value in relevant_values):
            continue

        records.append({
            "row_number": row_number,
            "employee_id": normalize_employee_id(get("employee_id")),
            "document": normalize_employee_id(get("document")),
            "first_name": clean_text(get("first_name")),
            "last_name": clean_text(get("last_name")),
            "contact_number": clean_text(get("contact_number")),
            "blood_type": clean_text(get("blood_type")),
            "vehicle_type": clean_text(get("vehicle_type")),
            "vehicle_brand": clean_text(get("vehicle_brand")),
            "vehicle_model": clean_text(get("vehicle_model")),
            "vehicle_color": clean_text(get("vehicle_color")),
            "vehicle_plate": clean_text(get("vehicle_plate")),
            "department": clean_text(get("department")),
            "role": clean_text(get("role")),
            "card_number": clean_text(get("card_number")),
            "secondary_card_number": clean_text(get("secondary_card_number")),
            "facility_code_wfm": clean_text(get("facility_code_wfm")),
            "facility_code_se": clean_text(get("facility_code_se")),
            "photo_and_data": clean_text(get("photo_and_data")),
            "eps": clean_text(get("eps")),
            "six_digit_card_number": clean_text(get("six_digit_card_number")),
            "first_replacement": clean_text(get("first_replacement")),
            "first_replacement_date": clean_text(get("first_replacement_date")),
            "second_replacement": clean_text(get("second_replacement")),
            "second_replacement_date": clean_text(get("second_replacement_date")),
        })
    return records


def _save_security_snapshot(rows, batch, source_file, source_sheet):
    batch.rows_received = len(rows)
    batch.save(update_fields=["rows_received"])

    if not rows:
        raise ImportServiceError("Security general base contains no data rows.")

    valid_rows = []
    rejected_rows = 0
    seen_documents = set()

    for row in rows:
        if not row["document"]:
            rejected_rows += 1
            _issue(batch, row["row_number"], row["employee_id"], "Missing DOCUMENTO.")
            continue
        if row["document"] in seen_documents:
            rejected_rows += 1
            _issue(
                batch,
                row["row_number"],
                row["employee_id"],
                f"Duplicate DOCUMENTO in this file: {row['document']}.",
            )
            continue
        seen_documents.add(row["document"])
        valid_rows.append(row)

    if not valid_rows:
        raise ImportServiceError(
            "No valid Security general base rows were found. Import aborted; the previous active snapshot was kept."
        )

    with transaction.atomic(using="default"):
        SecurityInfoRecord.objects.filter(is_active=True).update(is_active=False)
        SecurityInfoRecord.objects.bulk_create([
            SecurityInfoRecord(
                employee_id=row["employee_id"],
                document=row["document"],
                first_name=row["first_name"],
                last_name=row["last_name"],
                contact_number=row["contact_number"],
                blood_type=row["blood_type"],
                vehicle_type=row["vehicle_type"],
                vehicle_brand=row["vehicle_brand"],
                vehicle_model=row["vehicle_model"],
                vehicle_color=row["vehicle_color"],
                vehicle_plate=row["vehicle_plate"],
                department=row["department"],
                role=row["role"],
                card_number=row["card_number"],
                secondary_card_number=row["secondary_card_number"],
                facility_code_wfm=row["facility_code_wfm"],
                facility_code_se=row["facility_code_se"],
                photo_and_data=row["photo_and_data"],
                eps=row["eps"],
                six_digit_card_number=row["six_digit_card_number"],
                first_replacement=row["first_replacement"],
                first_replacement_date=row["first_replacement_date"],
                second_replacement=row["second_replacement"],
                second_replacement_date=row["second_replacement_date"],
                source_file=source_file,
                source_sheet=source_sheet,
                import_batch=batch,
            )
            for row in valid_rows
        ])

    return _finish_batch(batch, len(valid_rows), rejected_rows)


def _import_security_workbook(workbook, batch, source_file):
    candidate = find_security_sheet(workbook)
    if candidate is None:
        raise ImportServiceError("Security general base sheet was not found.")

    _, sheet, header_map, header_row = candidate
    rows = _read_security_rows(
        sheet.iter_rows(min_row=header_row + 1, values_only=True),
        header_map,
        header_row + 1,
    )
    return _save_security_snapshot(rows, batch, source_file, sheet.title)


def _read_csv_file(path):
    raw = Path(path).read_bytes()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        raise ImportServiceError("CSV encoding could not be read.")

    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    return list(csv.reader(io.StringIO(text), dialect))


def _import_security_csv(path, batch, source_file):
    csv_rows = _read_csv_file(path)
    candidate = find_security_csv_header(csv_rows)
    if candidate is None:
        raise ImportServiceError(
            "CSV format was not recognized as the Security general base. "
            "Required columns include DOCUMENTO, NOMBRE, APELLIDOS, DEPARTAMENTO, CARGO and No TARJETA."
        )

    header_map, header_index = candidate
    data_rows = csv_rows[header_index + 1:]
    rows = _read_security_rows(
        data_rows,
        header_map,
        header_index + 2,
    )
    return _save_security_snapshot(rows, batch, source_file, "CSV")


def _create_or_skip_batch(path, detected, file_name, force):
    digest = file_sha256(path)

    if not force and ImportBatch.objects.filter(
        file_hash=digest,
        source_type=detected,
        status__in=[ImportBatch.Status.SUCCESS, ImportBatch.Status.PARTIAL],
    ).exists():
        return ImportBatch.objects.create(
            file_name=file_name,
            file_hash=digest,
            source_type=detected,
            status=ImportBatch.Status.SKIPPED,
            finished_at=timezone.now(),
        ), True

    return ImportBatch.objects.create(
        file_name=file_name,
        file_hash=digest,
        source_type=detected,
    ), False


def _mark_batch_failed(batch, exc):
    batch.status = ImportBatch.Status.FAILED
    batch.finished_at = timezone.now()
    batch.save(update_fields=["status", "finished_at"])
    _issue(batch, None, "", str(exc))


def import_workbook(path, source_type="AUTO", force=False, source_file_name=None):
    """Import one supported source file.

    The historical public function name is kept for compatibility, but the
    importer now accepts both Excel workbooks and CSV files.
    """
    path = Path(path)
    if not path.exists():
        raise ImportServiceError(f"File not found: {path}")

    requested = (source_type or "AUTO").upper()
    file_name = source_file_name or path.name
    suffix = path.suffix.lower()

    if suffix == ".csv":
        csv_rows = _read_csv_file(path)
        security_candidate = find_security_csv_header(csv_rows)

        if requested == DataUpload.SourceType.AUTO:
            if security_candidate is None:
                raise ImportServiceError(
                    "CSV format was not recognized. CSV upload is currently supported for the Security general base."
                )
            detected = ImportBatch.SourceType.SECURITY_GENERAL
        elif requested == ImportBatch.SourceType.SECURITY_GENERAL:
            detected = requested
        else:
            raise ImportServiceError(
                "CSV upload is supported for Security general base. "
                "Use an Excel workbook for the tattoo or appearance approvals trackers."
            )

        batch, skipped = _create_or_skip_batch(path, detected, file_name, force)
        if skipped:
            return batch

        try:
            return _import_security_csv(path, batch, file_name)
        except Exception as exc:
            _mark_batch_failed(batch, exc)
            if isinstance(exc, ImportServiceError):
                raise
            raise ImportServiceError(str(exc)) from exc

    if suffix not in SUPPORTED_EXCEL_SUFFIXES:
        raise ImportServiceError(
            "Unsupported file format. Supported formats are .csv, .xlsx, .xlsm, .xltx and .xltm."
        )

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ImportServiceError(f"Excel file could not be opened: {exc}") from exc

    try:
        if requested == DataUpload.SourceType.AUTO:
            detected = detect_source_type(workbook)
        elif requested in {
            ImportBatch.SourceType.TATTOOS,
            ImportBatch.SourceType.APPEARANCE_APPROVALS,
            ImportBatch.SourceType.SECURITY_GENERAL,
        }:
            detected = requested
        else:
            raise ImportServiceError(f"Unsupported source type: {source_type}")

        batch, skipped = _create_or_skip_batch(path, detected, file_name, force)
        if skipped:
            return batch

        try:
            if detected == ImportBatch.SourceType.TATTOOS:
                return _import_tattoos(workbook, batch, file_name)
            if detected == ImportBatch.SourceType.APPEARANCE_APPROVALS:
                return _import_approvals(workbook, batch, file_name)
            return _import_security_workbook(workbook, batch, file_name)
        except Exception as exc:
            _mark_batch_failed(batch, exc)
            if isinstance(exc, ImportServiceError):
                raise
            raise ImportServiceError(str(exc)) from exc
    finally:
        workbook.close()


def process_data_upload(upload, force=False):
    upload.status = ImportBatch.Status.PROCESSING
    upload.error_message = ""
    upload.save(update_fields=["status", "error_message"])

    try:
        batch = import_workbook(
            upload.file.path,
            source_type=upload.source_type,
            force=force,
            source_file_name=Path(upload.file.name).name,
        )
        upload.import_batch = batch
        upload.detected_source_type = batch.source_type
        upload.status = batch.status
        upload.processed_at = timezone.now()
        upload.error_message = ""
        upload.save(
            update_fields=[
                "import_batch",
                "detected_source_type",
                "status",
                "processed_at",
                "error_message",
            ]
        )
        return batch
    except Exception as exc:
        upload.status = ImportBatch.Status.FAILED
        upload.processed_at = timezone.now()
        upload.error_message = str(exc)
        upload.save(update_fields=["status", "processed_at", "error_message"])
        raise
