import hashlib
import re
import shutil
from datetime import date, datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from appearance.models import HiBobEmployee, ImportBatch, ImportIssue, TattooRecord
from appearance.services import normalize_employee_id


HEADER_ALIASES = {
    "id": "employee_id",
    "name": "source_name",
    "tattoos": "tattoo_details",
    "date": "source_date",
    "shouldbecovered": "should_be_covered",
    "connotation": "connotation",
    "where": "location",
}


def normalize_header(value):
    value = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9]", "", value)


def normalize_bool(value):
    if value is None or str(value).strip() == "":
        return None
    normalized = str(value).strip().lower()
    if normalized in {"yes", "y", "si", "sí", "true", "1"}:
        return True
    if normalized in {"no", "n", "false", "0"}:
        return False
    return None


def normalize_date(value, workbook_epoch):
    if value in {None, ""}:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, workbook_epoch)
            return converted.date() if isinstance(converted, datetime) else converted
        except (ValueError, TypeError):
            return None

    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def file_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Command(BaseCommand):
    help = "Import and standardize tattoo Excel files from data/inbox."

    def add_arguments(self, parser):
        parser.add_argument("--file", help="Optional specific xlsx file to import.")
        parser.add_argument("--force", action="store_true", help="Re-import a file even if its hash was already processed.")

    def handle(self, *args, **options):
        inbox = Path(settings.DATA_INBOX)
        processed = Path(settings.DATA_PROCESSED)
        rejected = Path(settings.DATA_REJECTED)
        inbox.mkdir(parents=True, exist_ok=True)
        processed.mkdir(parents=True, exist_ok=True)
        rejected.mkdir(parents=True, exist_ok=True)

        files = [Path(options["file"])] if options.get("file") else sorted(inbox.glob("*.xlsx"))
        if not files:
            self.stdout.write(self.style.WARNING("No .xlsx files found in data/inbox."))
            return

        for path in files:
            self._import_file(path, processed, rejected, options["force"])

    def _import_file(self, path, processed_dir, rejected_dir, force):
        if not path.exists():
            raise CommandError(f"File not found: {path}")

        file_hash = file_sha256(path)
        if not force and ImportBatch.objects.filter(file_hash=file_hash, status__in=[ImportBatch.Status.SUCCESS, ImportBatch.Status.PARTIAL]).exists():
            self.stdout.write(self.style.WARNING(f"Skipping already imported file: {path.name}"))
            return

        batch = ImportBatch.objects.create(file_name=path.name, file_hash=file_hash, source_type="TATTOOS")

        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
            sheet, header_map, header_row = self._find_best_sheet(workbook)
            rows = self._read_rows(sheet, header_map, header_row, workbook.epoch)
            batch.rows_received = len(rows)
            batch.save(update_fields=["rows_received"])

            employee_ids = [row["employee_id"] for row in rows if row["employee_id"]]
            known_ids = set(
                HiBobEmployee.objects.filter(employee_id__in=employee_ids)
                .values_list("employee_id", flat=True)
                .distinct()
            )

            valid_rows = []
            for row in rows:
                if not row["employee_id"]:
                    self._issue(batch, row["row_number"], "", "Missing Employee ID.")
                    continue
                if row["employee_id"] not in known_ids:
                    self._issue(batch, row["row_number"], row["employee_id"], "Employee ID was not found in HiBob.")
                    continue
                valid_rows.append(row)

            with transaction.atomic():
                touched_ids = {row["employee_id"] for row in valid_rows}
                TattooRecord.objects.filter(employee_id__in=touched_ids, is_active=True).update(is_active=False)

                TattooRecord.objects.bulk_create([
                    TattooRecord(
                        employee_id=row["employee_id"],
                        source_name=row["source_name"],
                        tattoo_details=row["tattoo_details"],
                        should_be_covered=row["should_be_covered"],
                        connotation=row["connotation"],
                        location=row["location"],
                        source_date=row["source_date"],
                        source_sheet=sheet.title,
                        source_file=path.name,
                        import_batch=batch,
                    )
                    for row in valid_rows
                ])

            batch.rows_imported = len(valid_rows)
            batch.rows_rejected = batch.issues.count()
            batch.status = ImportBatch.Status.PARTIAL if batch.rows_rejected else ImportBatch.Status.SUCCESS
            batch.finished_at = timezone.now()
            batch.save(update_fields=["rows_imported", "rows_rejected", "status", "finished_at"])

            target = processed_dir / path.name
            if path.parent.resolve() == Path(settings.DATA_INBOX).resolve():
                if target.exists():
                    target.unlink()
                shutil.move(str(path), str(target))

            self.stdout.write(self.style.SUCCESS(
                f"{path.name}: {batch.rows_imported} imported, {batch.rows_rejected} rejected, status={batch.status}."
            ))
        except Exception as exc:
            batch.status = ImportBatch.Status.FAILED
            batch.finished_at = timezone.now()
            batch.save(update_fields=["status", "finished_at"])
            self._issue(batch, None, "", str(exc))
            if path.parent.resolve() == Path(settings.DATA_INBOX).resolve():
                target = rejected_dir / path.name
                if target.exists():
                    target.unlink()
                shutil.move(str(path), str(target))
            raise CommandError(f"Failed to import {path.name}: {exc}") from exc

    def _find_best_sheet(self, workbook):
        candidates = []
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10), values_only=True), start=1):
                normalized = [normalize_header(value) for value in row]
                if "id" in normalized and "tattoos" in normalized and "shouldbecovered" in normalized:
                    header_map = {}
                    for index, header in enumerate(normalized):
                        mapped = HEADER_ALIASES.get(header)
                        if mapped:
                            header_map[mapped] = index
                    candidates.append((sheet.max_row, sheet, header_map, row_number))
                    break
        if not candidates:
            raise ValueError("No worksheet with ID, TATTOOS and SHOULD BE COVERED? columns was found.")
        _, sheet, header_map, header_row = max(candidates, key=lambda item: item[0])
        return sheet, header_map, header_row

    def _read_rows(self, sheet, header_map, header_row, workbook_epoch):
        records = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            def get(field):
                index = header_map.get(field)
                return values[index] if index is not None and index < len(values) else None

            employee_id = normalize_employee_id(get("employee_id"))
            if not employee_id and all(value in {None, ""} for value in values):
                continue

            records.append({
                "row_number": row_number,
                "employee_id": employee_id,
                "source_name": str(get("source_name") or "").strip(),
                "tattoo_details": str(get("tattoo_details") or "").strip(),
                "should_be_covered": normalize_bool(get("should_be_covered")),
                "connotation": str(get("connotation") or "").strip(),
                "location": str(get("location") or "").strip(),
                "source_date": normalize_date(get("source_date"), workbook_epoch),
            })
        return records

    @staticmethod
    def _issue(batch, row_number, employee_id, message):
        ImportIssue.objects.create(
            batch=batch,
            row_number=row_number,
            employee_id=employee_id,
            message=message,
        )
